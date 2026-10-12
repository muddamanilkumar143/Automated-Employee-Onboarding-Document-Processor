"""Automated Employee Onboarding Document Processor."""
import argparse
import csv
import logging
import re
from datetime import datetime, date
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

BASE_DIR = Path(__file__).resolve().parent
REQUIRED_FIELDS = ["employee_id", "name", "department", "role", "joining_date", "manager"]
FIELD_LABELS = {
    "employee_id": "Employee ID", "name": "Employee Name",
    "department": "Department", "role": "Role",
    "joining_date": "Joining Date", "manager": "Reporting Manager",
}
LOGGER = logging.getLogger("onboarding")


def setup_logging(log_dir: Path) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    log_file = log_dir / "onboarding.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
        handlers=[logging.FileHandler(log_file, encoding="utf-8"), logging.StreamHandler()],
        force=True,
    )
    return log_file


def normalize_header(header):
    return str(header or "").strip().lower().replace(" ", "_")

def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise ValueError("Joining date must use YYYY-MM-DD, DD-MM-YYYY, DD/MM/YYYY, or MM/DD/YYYY")


def read_employee_records(input_path: Path):
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")
    suffix = input_path.suffix.lower()

    if suffix == ".csv":
        with input_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if not reader.fieldnames:
                raise ValueError(f"CSV file is empty or has no headers: {input_path}")
            headers = [normalize_header(h) for h in reader.fieldnames]
            missing = [field for field in REQUIRED_FIELDS if field not in headers]
            if missing:
                raise ValueError("Missing required CSV columns: " + ", ".join(missing))
            return [{normalize_header(k): v for k, v in row.items() if k is not None} for row in reader]

    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("Excel support requires openpyxl. Run: python -m pip install openpyxl") from exc
        workbook = load_workbook(input_path, read_only=True, data_only=True)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            header_row = next(iterator, None)
            if not header_row:
                raise ValueError(f"Excel file is empty or has no headers: {input_path}")
            headers = [normalize_header(h) for h in header_row]
            missing = [field for field in REQUIRED_FIELDS if field not in headers]
            if missing:
                raise ValueError("Missing required Excel columns: " + ", ".join(missing))
            records = []
            for row in iterator:
                if any(v is not None and str(v).strip() for v in row):
                    records.append(dict(zip(headers, row)))
            return records
        finally:
            workbook.close()

    raise ValueError("Unsupported input format. Use .csv, .xlsx, or .xlsm.")


def validate_employee(record, seen_ids):
    employee = {
        field: str(record.get(field) if record.get(field) is not None else "").strip()
        for field in REQUIRED_FIELDS
    }
    errors = []
    for field in REQUIRED_FIELDS:
        if not employee[field]:
            errors.append(f"{FIELD_LABELS[field]} is required")

    employee_id = employee["employee_id"]
    if employee_id:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", employee_id):
            errors.append("Employee ID may contain only letters, numbers, underscores, and hyphens")
        normalized_id = employee_id.casefold()
        if normalized_id in seen_ids:
            errors.append(f"Duplicate Employee ID: {employee_id}")
        else:
            seen_ids.add(normalized_id)

    if employee["joining_date"]:
        try:
            employee["joining_date"] = parse_date(record.get("joining_date")).strftime("%Y-%m-%d")
        except (ValueError, TypeError):
            errors.append("Joining date is invalid")

    return (None, errors) if errors else (employee, [])


def safe_folder_name(value):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_") or "employee"


def render_template(template_text, employee):
    values = dict(employee)
    values["processed_date"] = date.today().strftime("%Y-%m-%d")
    output = template_text
    for key, value in values.items():
        output = output.replace("{{" + key + "}}", str(value))
    unresolved = re.findall(r"\{\{([A-Za-z0-9_]+)\}\}", output)
    if unresolved:
        raise ValueError("Unknown template placeholder(s): " + ", ".join(sorted(set(unresolved))))
    return output


def save_as_pdf(document_text, pdf_path: Path, employee):
    """Create a formatted A4 PDF using ReportLab."""
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "OnboardingTitle", parent=styles["Title"], alignment=TA_CENTER,
        textColor=colors.HexColor("#17365D"), fontSize=18, leading=22, spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "OnboardingSubtitle", parent=styles["Normal"], alignment=TA_CENTER,
        textColor=colors.HexColor("#555555"), fontSize=9, spaceAfter=18,
    )
    label_style = ParagraphStyle("FieldLabel", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10, leading=15)
    value_style = ParagraphStyle("FieldValue", parent=styles["Normal"], fontSize=10, leading=15)
    body_style = ParagraphStyle("Checklist", parent=styles["BodyText"], fontSize=9, leading=14, spaceAfter=5)

    story = [
        Paragraph("EMPLOYEE ONBOARDING RECORD", title_style),
        Paragraph("Generated by the Employee Onboarding Processor", subtitle_style),
        Spacer(1, 5 * mm),
    ]
    rows = []
    for field in REQUIRED_FIELDS:
        value = str(employee.get(field, ""))
        value = value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        rows.append([Paragraph(FIELD_LABELS[field], label_style), Paragraph(value, value_style)])
    table = Table(rows, colWidths=[48 * mm, 105 * mm], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF0F7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C7D2E0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([table, Spacer(1, 10 * mm), Paragraph("ONBOARDING CHECKLIST", styles["Heading2"])])
    for item in [
        "[ ] Verify identity and employment documents",
        "[ ] Complete company policy orientation",
        "[ ] Set up company accounts and required access",
        "[ ] Review role responsibilities with manager",
        "[ ] Confirm first-week schedule",
    ]:
        story.append(Paragraph(item, body_style))
    story.extend([
        Spacer(1, 8 * mm),
        Paragraph(f"Processed date: {date.today().strftime('%Y-%m-%d')}", styles["Italic"]),
        Spacer(1, 4 * mm),
        Paragraph("Note: Verify the information before using this document for actual onboarding.", styles["Italic"]),
    ])
    SimpleDocTemplate(
        str(pdf_path), pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=f"Onboarding - {employee.get('employee_id', '')}",
        author="Employee Onboarding Processor",
    ).build(story)


def write_summary(summary, reports_dir: Path):
    reports_dir.mkdir(parents=True, exist_ok=True)
    summary_file = reports_dir / "processing_summary.csv"
    fields = ["row_number", "employee_id", "name", "status", "message", "txt_file", "pdf_file"]
    with summary_file.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(summary)
    return summary_file


def process_employees(records, template_text, output_dir: Path, reports_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    seen_ids, summary = set(), []

    for row_number, record in enumerate(records, start=2):
        employee_id = str(record.get("employee_id") or "").strip()
        name = str(record.get("name") or "").strip()
        employee, errors = validate_employee(record, seen_ids)
        if errors:
            message = "; ".join(errors)
            LOGGER.warning("Row %s rejected (ID=%s): %s", row_number, employee_id or "UNKNOWN", message)
            summary.append({
                "row_number": row_number, "employee_id": employee_id, "name": name,
                "status": "FAILED", "message": message, "txt_file": "", "pdf_file": "",
            })
            continue

        folder = output_dir / safe_folder_name(f"{employee['employee_id']}_{employee['name']}")
        txt_path, pdf_path = folder / "onboarding.txt", folder / "onboarding.pdf"
        try:
            folder.mkdir(parents=True, exist_ok=True)
            document = render_template(template_text, employee)
            txt_path.write_text(document, encoding="utf-8")
            save_as_pdf(document, pdf_path, employee)
            LOGGER.info("Generated TXT and PDF for %s", employee["employee_id"])
            summary.append({
                "row_number": row_number, "employee_id": employee["employee_id"],
                "name": employee["name"], "status": "SUCCESS",
                "message": "TXT and PDF generated",
                "txt_file": str(txt_path.relative_to(BASE_DIR)),
                "pdf_file": str(pdf_path.relative_to(BASE_DIR)),
            })
        except (OSError, ValueError, KeyError) as exc:
            LOGGER.exception("Document generation failed for %s", employee_id)
            summary.append({
                "row_number": row_number, "employee_id": employee_id, "name": name,
                "status": "FAILED", "message": f"{type(exc).__name__}: {exc}",
                "txt_file": "", "pdf_file": "",
            })

    summary_file = write_summary(summary, reports_dir)
    success = sum(row["status"] == "SUCCESS" for row in summary)
    failed = sum(row["status"] == "FAILED" for row in summary)
    LOGGER.info("Finished: total=%d success=%d failed=%d", len(summary), success, failed)
    print("\\n===== PROCESSING SUMMARY =====")
    print(f"Total records: {len(summary)}")
    print(f"Successful: {success}")
    print(f"Failed: {failed}")
    print(f"Summary: {summary_file}")
    return {"total": len(summary), "success": success, "failed": failed}


def main():
    parser = argparse.ArgumentParser(description="Generate employee onboarding TXT and PDF documents.")
    parser.add_argument("--input", default=str(BASE_DIR / "input" / "employees.csv"))
    parser.add_argument("--template", default=str(BASE_DIR / "templates" / "onboarding_template.txt"))
    parser.add_argument("--output", default=str(BASE_DIR / "output"))
    parser.add_argument("--reports", default=str(BASE_DIR / "reports"))
    parser.add_argument("--logs", default=str(BASE_DIR / "logs"))
    args = parser.parse_args()
    log_file = setup_logging(Path(args.logs))
    LOGGER.info("Starting onboarding processor")
    try:
        template_path = Path(args.template)
        if not template_path.exists():
            raise FileNotFoundError(f"Template file not found: {template_path}")
        records = read_employee_records(Path(args.input))
        template = template_path.read_text(encoding="utf-8")
        process_employees(records, template, Path(args.output), Path(args.reports))
        print(f"Log: {log_file}")
    except (OSError, ValueError, csv.Error, RuntimeError) as exc:
        LOGGER.exception("Processing stopped: %s", exc)
        print(f"ERROR: {exc}")


if __name__ == "__main__":
    main()
